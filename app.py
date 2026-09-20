import sqlite3
import os
import io
import re
from datetime import date
from functools import wraps

from markupsafe import Markup

from flask import (
    Flask, request, redirect, url_for, render_template,
    make_response, send_file, abort, flash
)

from config import DB_PATH
from session_utils import Usuario, crear_cookie, leer_cookie

app = Flask(__name__)
app.secret_key = "flash-messages-key"

COOKIE_NAME = "sesion"


# --------------------------------------------------------------------------
#  Formato de la descripcion de una propuesta
# --------------------------------------------------------------------------
@app.template_filter("formato")
def formato(texto):
    """Da formato basico a la descripcion de una propuesta: **negrita** y
    saltos de linea, para mostrarla prolija en el panel del admin."""
    html = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", texto or "")  # **negrita**
    html = html.replace("\n", "<br>")                                         # saltos de linea
    return Markup(html)   # se marca como HTML para renderizar el formato


# --------------------------------------------------------------------------
#  Helpers de base de datos
# --------------------------------------------------------------------------
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def cierre_vencido(fecha_cierre):
    """True si la fecha de cierre (formato YYYY-MM-DD) ya paso respecto de hoy."""
    if not fecha_cierre:
        return False
    try:
        return date.fromisoformat(fecha_cierre) < date.today()
    except ValueError:
        return False


# --------------------------------------------------------------------------
#  Manejo de sesion (cookie personalizada pickle + HMAC)
# --------------------------------------------------------------------------
def usuario_actual():
    """Devuelve el objeto Usuario de la cookie, o None si no hay sesion valida."""
    valor = request.cookies.get(COOKIE_NAME)
    if not valor:
        return None
    try:
        return leer_cookie(valor)   # <-- Vuln 4: deserializacion insegura
    except Exception:
        return None


def login_requerido(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if usuario_actual() is None:
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return wrapper


def admin_requerido(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        u = usuario_actual()
        if u is None:
            return redirect(url_for("login"))
        if not u.es_admin():
            abort(403)
        return f(*args, **kwargs)
    return wrapper


@app.context_processor
def inyectar_usuario():
    """Hace disponible 'usuario' en todas las plantillas."""
    return {"usuario": usuario_actual()}


# --------------------------------------------------------------------------
#  Rutas publicas
# --------------------------------------------------------------------------
@app.route("/")
def index():
    u = usuario_actual()

    # Sin sesion iniciada solo se muestran las licitaciones abiertas.
    # Con sesion iniciada se puede filtrar entre abiertas y cerradas.
    if u is None:
        filtro = "abierta"
    else:
        filtro = request.args.get("estado", "todas")

    db = get_db()
    if filtro in ("abierta", "cerrada"):
        licitaciones = db.execute(
            "SELECT * FROM licitaciones WHERE estado = ? ORDER BY estado, fecha_cierre",
            (filtro,),
        ).fetchall()
    else:
        licitaciones = db.execute(
            "SELECT * FROM licitaciones ORDER BY estado, fecha_cierre"
        ).fetchall()

    # Licitaciones en las que el usuario actual ya presento una propuesta
    presentadas = set()
    if u is not None and not u.es_admin():
        filas = db.execute(
            "SELECT licitacion_id FROM propuestas WHERE usuario_id = ?", (u.id,)
        ).fetchall()
        presentadas = {f["licitacion_id"] for f in filas}

    db.close()
    vencidas = {l["id"] for l in licitaciones if cierre_vencido(l["fecha_cierre"])}
    return render_template("index.html", licitaciones=licitaciones,
                           filtro=filtro, presentadas=presentadas,
                           vencidas=vencidas)


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "")
        password = request.form.get("password", "")
        db = get_db()
        fila = db.execute(
            "SELECT * FROM usuarios WHERE username = ? AND password = ?",
            (username, password),
        ).fetchone()
        db.close()

        if fila is None:
            flash("Usuario o contrasena incorrectos.", "error")
            return render_template("login.html")

        usuario = Usuario(
            id=fila["id"],
            username=fila["username"],
            rol=fila["rol"],
            nombre_empresa=fila["nombre_empresa"],
        )
        resp = make_response(redirect(url_for("index")))
        # -------------------------------------------------------------
        #  Vuln 2 (Broken Auth / Session): cookie SIN HttpOnly,
        #  sin Secure ni SameSite. JavaScript puede leerla con
        #  document.cookie (necesario para el robo via XSS).
        # -------------------------------------------------------------
        resp.set_cookie(
            COOKIE_NAME,
            crear_cookie(usuario),
            httponly=False,   # <-- deliberadamente False
            secure=False,
            samesite=None,
        )
        return resp

    return render_template("login.html")


@app.route("/logout")
def logout():
    resp = make_response(redirect(url_for("index")))
    resp.delete_cookie(COOKIE_NAME)
    return resp


# --------------------------------------------------------------------------
#  Propuestas (usuarios logueados)
# --------------------------------------------------------------------------
@app.route("/propuesta/<int:licitacion_id>", methods=["GET", "POST"])
@login_requerido
def propuesta(licitacion_id):
    u = usuario_actual()

    # Un administrador gestiona licitaciones; no participa como proveedor.
    if u.es_admin():
        flash("Los administradores no pueden presentar propuestas.", "error")
        return redirect(url_for("admin_licitacion", id=licitacion_id))

    db = get_db()
    lic = db.execute(
        "SELECT * FROM licitaciones WHERE id = ?", (licitacion_id,)
    ).fetchone()

    if lic is None:
        db.close()
        abort(404)

    # No se puede presentar si la fecha de cierre ya paso.
    if cierre_vencido(lic["fecha_cierre"]):
        db.close()
        flash("La fecha de cierre de esta licitacion ya paso. No se admiten nuevas propuestas.", "error")
        return redirect(url_for("index"))

    # No se puede presentar mas de una propuesta para la misma licitacion.
    ya_presentada = db.execute(
        "SELECT 1 FROM propuestas WHERE licitacion_id = ? AND usuario_id = ?",
        (licitacion_id, u.id),
    ).fetchone()
    if ya_presentada:
        db.close()
        flash("Ya presentaste una propuesta para esta licitacion.", "error")
        return redirect(url_for("mis_propuestas"))

    if request.method == "POST":
        razon_social = request.form.get("razon_social", "").strip()
        cuit = request.form.get("cuit", "").strip()
        monto_raw = request.form.get("monto", "").strip()
        descripcion = request.form.get("descripcion", "")   # <-- NO sanitizado
        email_contacto = request.form.get("email_contacto", "").strip()

        errores = []

        # --- Validacion de monto
        try:
            monto = float(monto_raw)
            if monto <= 0:
                errores.append("El monto debe ser un numero positivo.")
        except ValueError:
            errores.append("El monto debe ser un numero valido.")
            monto = None

        # --- Validacion de CUIT
        import re
        if not re.fullmatch(r"\d{2}-\d{8}-\d", cuit):
            errores.append("El CUIT debe tener el formato XX-XXXXXXXX-X.")

        if not razon_social:
            errores.append("La razon social es obligatoria.")

        # --- Validacion del email de contacto
        if not email_contacto:
            errores.append("El email de contacto es obligatorio.")
        elif not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email_contacto):
            errores.append("El email de contacto no tiene un formato valido.")


        if errores:
            db.close()
            return render_template(
                "propuesta.html", lic=lic, errores=errores,
                form=request.form,
            )

        db.execute(
            """INSERT INTO propuestas
               (licitacion_id, usuario_id, razon_social, cuit, monto,
                descripcion, email_contacto, estado)
               VALUES (?,?,?,?,?,?,?, 'pendiente')""",
            (licitacion_id, u.id, razon_social, cuit, monto,
             descripcion, email_contacto),
        )
        db.commit()
        db.close()
        flash("Propuesta enviada correctamente.", "ok")
        return redirect(url_for("mis_propuestas"))

    db.close()
    return render_template("propuesta.html", lic=lic, errores=None, form={})


@app.route("/mis-propuestas")
@login_requerido
def mis_propuestas():
    u = usuario_actual()

    # El admin no presenta propuestas; gestiona desde el panel.
    if u.es_admin():
        return redirect(url_for("admin_panel"))

    db = get_db()
    propuestas = db.execute(
        """SELECT p.*, l.titulo AS lic_titulo
           FROM propuestas p
           JOIN licitaciones l ON l.id = p.licitacion_id
           WHERE p.usuario_id = ?
           ORDER BY p.id DESC""",
        (u.id,),
    ).fetchall()
    db.close()
    return render_template("seguimiento.html", propuestas=propuestas)


@app.route("/propuesta/<int:propuesta_id>/eliminar", methods=["POST"])
@login_requerido
def eliminar_propuesta(propuesta_id):
    u = usuario_actual()
    db = get_db()
    prop = db.execute(
        "SELECT * FROM propuestas WHERE id = ?", (propuesta_id,)
    ).fetchone()

    # Solo el dueno de la propuesta puede eliminarla.
    if prop is None or prop["usuario_id"] != u.id:
        db.close()
        abort(404)

    db.execute("DELETE FROM propuestas WHERE id = ?", (propuesta_id,))
    db.commit()
    db.close()
    flash("Propuesta eliminada.", "ok")
    return redirect(url_for("mis_propuestas"))


@app.route("/perfil", methods=["GET", "POST"])
@login_requerido
def perfil():
    u = usuario_actual()
    db = get_db()

    if request.method == "POST":
        nombre_empresa = request.form.get("nombre_empresa", "").strip()
        db.execute(
            "UPDATE usuarios SET nombre_empresa = ? WHERE id = ?",
            (nombre_empresa, u.id),
        )
        db.commit()
        db.close()
        flash("Perfil actualizado. Volve a iniciar sesion para ver los cambios.", "ok")
        return redirect(url_for("perfil"))

    fila = db.execute("SELECT * FROM usuarios WHERE id = ?", (u.id,)).fetchone()
    db.close()
    return render_template("perfil.html", perfil=fila)


# --------------------------------------------------------------------------
#  Panel de administracion (solo admin)
# --------------------------------------------------------------------------
@app.route("/admin")
@admin_requerido
def admin_panel():
    db = get_db()
    licitaciones = db.execute(
        """SELECT l.*, COUNT(p.id) AS n_propuestas
           FROM licitaciones l
           LEFT JOIN propuestas p ON p.licitacion_id = l.id
           GROUP BY l.id
           ORDER BY l.id"""
    ).fetchall()
    db.close()
    # Reportes disponibles
    reportes = []
    if os.path.isdir("reportes"):
        reportes = sorted(os.listdir("reportes"))
    return render_template("admin_panel.html",
                           licitaciones=licitaciones, reportes=reportes)


@app.route("/admin/licitacion/<int:id>")
@admin_requerido
def admin_licitacion(id):
    db = get_db()
    lic = db.execute("SELECT * FROM licitaciones WHERE id = ?", (id,)).fetchone()
    if lic is None:
        db.close()
        abort(404)
    propuestas = db.execute(
        """SELECT p.*, u.username
           FROM propuestas p
           JOIN usuarios u ON u.id = p.usuario_id
           WHERE p.licitacion_id = ?
           ORDER BY p.monto ASC""",
        (id,),
    ).fetchall()
    db.close()
    return render_template("admin_licitacion.html", lic=lic, propuestas=propuestas)


@app.route("/admin/propuesta/<int:propuesta_id>/estado", methods=["POST"])
@admin_requerido
def admin_estado_propuesta(propuesta_id):
    accion = request.form.get("accion", "")
    # 'no_seleccionada' es un estado AUTOMATICO (no es una accion elegible).
    estados = {
        "aprobar": "aprobada",
        "rechazar": "rechazada",
        "seleccionar": "seleccionada",
        "pendiente": "pendiente",
    }
    nuevo_estado = estados.get(accion)

    db = get_db()
    prop = db.execute(
        "SELECT * FROM propuestas WHERE id = ?", (propuesta_id,)
    ).fetchone()
    if prop is None or nuevo_estado is None:
        db.close()
        abort(404)

    lic_id = prop["licitacion_id"]

    if nuevo_estado == "seleccionada":
        # Seleccionar = adjudicar: esta propuesta queda seleccionada, TODAS
        # las demas pasan a 'no_seleccionada' y la licitacion se cierra.
        db.execute(
            "UPDATE propuestas SET estado = 'no_seleccionada' "
            "WHERE licitacion_id = ? AND id <> ?",
            (lic_id, propuesta_id),
        )
        db.execute(
            "UPDATE propuestas SET estado = 'seleccionada' WHERE id = ?",
            (propuesta_id,),
        )
        db.execute(
            "UPDATE licitaciones SET estado = 'cerrada' WHERE id = ?",
            (lic_id,),
        )
        mensaje = ("Propuesta seleccionada. La licitacion se cerro y las demas "
                   "quedaron como 'no seleccionada'.")
    else:
        db.execute(
            "UPDATE propuestas SET estado = ? WHERE id = ?",
            (nuevo_estado, propuesta_id),
        )
        mensaje = "Propuesta actualizada."

    db.commit()
    db.close()
    flash(mensaje, "ok")
    return redirect(url_for("admin_licitacion", id=lic_id))


@app.route("/admin/licitacion/<int:id>/exportar")
@admin_requerido
def exportar_propuestas(id):
    """Genera un Excel (.xlsx) EN MEMORIA con las propuestas actuales de la
    licitacion y lo devuelve como descarga. No se guarda nada en disco:
    siempre refleja la informacion vigente en la base de datos."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    db = get_db()
    lic = db.execute("SELECT * FROM licitaciones WHERE id = ?", (id,)).fetchone()
    if lic is None:
        db.close()
        abort(404)
    propuestas = db.execute(
        """SELECT p.*, u.username
           FROM propuestas p
           JOIN usuarios u ON u.id = p.usuario_id
           WHERE p.licitacion_id = ?
           ORDER BY p.monto ASC""",
        (id,),
    ).fetchall()
    db.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Propuestas"

    encabezados = ["Proveedor", "Razon social", "CUIT", "Contacto", "Monto",
                   "Descripcion", "Estado"]
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True)

    for p in propuestas:
        ws.append([
            p["username"], p["razon_social"], p["cuit"], p["email_contacto"],
            p["monto"], p["descripcion"], p["estado"].replace("_", " "),
        ])

    # Formato de moneda en la columna Monto (columna E).
    for fila in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        fila[0].number_format = '#,##0'

    # Anchos de columna comodos.
    anchos = [16, 26, 16, 28, 14, 40, 16]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[chr(64 + i)].width = ancho

    # Se escribe el workbook a un buffer en memoria (no toca el disco).
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre = f"propuestas_licitacion_{id}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nombre,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/admin/nueva-licitacion", methods=["GET", "POST"])
@admin_requerido
def admin_nueva_licitacion():
    if request.method == "POST":
        titulo = request.form.get("titulo", "").strip()
        descripcion = request.form.get("descripcion", "").strip()
        estado = request.form.get("estado", "abierta")
        fecha_cierre = request.form.get("fecha_cierre", "").strip()

        if not titulo:
            return render_template("admin_nueva_licitacion.html",
                                   error="El titulo es obligatorio.",
                                   form=request.form)
        db = get_db()
        db.execute(
            """INSERT INTO licitaciones (titulo, descripcion, estado, fecha_cierre)
               VALUES (?,?,?,?)""",
            (titulo, descripcion, estado, fecha_cierre),
        )
        db.commit()
        db.close()
        flash("Licitacion creada.", "ok")
        return redirect(url_for("admin_panel"))

    return render_template("admin_nueva_licitacion.html", error=None, form={})


# --------------------------------------------------------------------------
#  Descarga de reportes (solo admin)  --  Vuln 3: Path Traversal / LFI
# --------------------------------------------------------------------------
@app.route("/admin/descargar")
@admin_requerido
def descargar():
    # -------------------------------------------------------------
    #  Vuln 3: se concatena el nombre recibido SIN sanitizar.
    #  ?file=../config.py  ->  lee config.py fuera de reportes/.
    # -------------------------------------------------------------
    nombre = request.args.get("file", "")
    ruta = os.path.join("reportes", nombre)   # sin sanitizar
    try:
        return send_file(os.path.abspath(ruta))
    except Exception:
        abort(404)


# --------------------------------------------------------------------------
#  Manejo de errores
# --------------------------------------------------------------------------
@app.errorhandler(403)
def prohibido(e):
    return render_template("base.html", contenido_403=True), 403


if __name__ == "__main__":
    if not os.path.exists(DB_PATH):
        print("[!] No existe database.db. Ejecuta primero:  python init_db.py")
    print("=" * 60)
    print("  Licitaciones Ficticio S.A.")
    print("  http://127.0.0.1:5000")
    print("=" * 60)
    app.run(host="127.0.0.1", port=5000, debug=True)
